"""Build Rehearsing-Sovereignty.docx from ARTICLE.md.

Usage: python build_docx.py
Figures: a line containing a bracketed figure note like
  [FIGURE: ... a1_classifier.png ...]  or  [FIGURE: ... usher_fig2.png ...]
embeds that image from analyses/ at that point, with the bracket text
(minus the filename) as the caption. A bracketed note naming no existing
file is kept as an italic placeholder paragraph.
"""
import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "ARTICLE.md")
OUT = os.path.join(HERE, "Rehearsing-Sovereignty.docx")
FIGDIR = os.path.join(HERE, "analyses")

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITAL_RE = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
FIG_RE = re.compile(r"\[FIGURE:([^\]]+)\]")


def add_runs(par, text):
    """Add text to a paragraph, honoring **bold** and *italic* spans."""
    pos = 0
    tokens = []
    for m in BOLD_RE.finditer(text):
        tokens.append((m.start(), m.end(), m.group(1), "b"))
    for m in ITAL_RE.finditer(text):
        if not any(s <= m.start() < e for s, e, _, _ in tokens):
            tokens.append((m.start(), m.end(), m.group(1), "i"))
    for start, end, inner, kind in sorted(tokens):
        if start > pos:
            par.add_run(text[pos:start])
        run = par.add_run(inner)
        if kind == "b":
            run.bold = True
        else:
            run.italic = True
        pos = end
    if pos < len(text):
        par.add_run(text[pos:])


def find_figure(note):
    """Return (paths, caption) for every known image file named in the note."""
    paths = []
    caption = note
    for name in re.findall(r"[\w\-]+\.(?:png|jpg|jpeg)", note):
        path = os.path.join(FIGDIR, name)
        if os.path.exists(path):
            paths.append(path)
            caption = caption.replace(name, "")
    caption = re.sub(r"\s{2,}", " ", caption).strip(" .,:")
    return paths, caption


def build():
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)

    lines = open(SRC, encoding="utf-8").read().splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        fig = FIG_RE.search(stripped)
        if fig:
            before = stripped[: fig.start()].strip()
            if before:
                add_runs(doc.add_paragraph(), before)
            paths, caption = find_figure(fig.group(1))
            if paths:
                width = Inches(6.0) if len(paths) == 1 else Inches(2.0)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for path in paths:
                    p.add_run().add_picture(path, width=width)
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(caption)
                run.italic = True
                run.font.size = Pt(9)
            else:
                p = doc.add_paragraph()
                run = p.add_run("[Figure: " + caption + "]")
                run.italic = True
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level, 4))
        elif stripped.startswith(("- ", "* ")):
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            add_runs(
                doc.add_paragraph(style="List Number"),
                re.sub(r"^\d+\.\s+", "", stripped),
            )
        else:
            add_runs(doc.add_paragraph(), stripped)

    doc.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    build()
